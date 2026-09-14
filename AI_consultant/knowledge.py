import io
import os
import json
import re
import base64
import hashlib
import threading
import logging
import mimetypes
from pathlib import Path
from typing import List, Optional, Dict, Any
from datetime import datetime, timezone

import httpx
from urllib.parse import urlparse
from sqlalchemy import text as sa_text, select, delete as sa_delete

from models import KnowledgeBase, KnowledgeDocument, KnowledgeChunk, Base
from db import _engine as get_db_engine, get_db, get_redis

logger = logging.getLogger(__name__)

# ── 文本分块配置 ────────────────────────────────────────────
DEFAULT_CHUNK_SIZE = 1024
DEFAULT_CHUNK_OVERLAP = 50
DEFAULT_CHUNK_SEPARATOR = "\n\n"
MAX_CHUNKS_PER_DOC = 10000

_HTTP_CLIENT: Optional[httpx.Client] = None
_JSON_MEDIA_TYPE = "application/json"


def _utcnow() -> datetime:
    return datetime.now(timezone.utc).replace(tzinfo=None)


def _or_default(value, fallback):
    return value or fallback


def _get_http_client() -> httpx.Client:
    global _HTTP_CLIENT
    if _HTTP_CLIENT is None:
        _HTTP_CLIENT = httpx.Client(timeout=120, headers={"Content-Type": _JSON_MEDIA_TYPE})
    return _HTTP_CLIENT


# ═══════════════════════════════════════════════════════════════
#  1. Embedding 客户端（基于 .env 配置，原始 HTTP 请求）
# ═══════════════════════════════════════════════════════════════
class EmbeddingClient:
    _cache_ttl = 3600
    _memory_cache: Dict[str, List[float]] = {}
    _lock = threading.Lock()

    def __init__(self, config):
        self.config = config

    @property
    def api_key(self) -> str:
        val = self.config.embedding_api_key or self.config.llm_api_key or ""
        return val.strip()

    @property
    def base_url(self) -> str:
        raw = self.config.embedding_base_url or self.config.llm_base_url or ""
        return raw.rstrip("/")

    @property
    def model(self) -> str:
        return self.config.embedding_model

    @property
    def dimension(self) -> int:
        return self.config.embedding_dimension or 1536

    def _cache_key(self, text: str) -> str:
        h = hashlib.md5(text.encode("utf-8"), usedforsecurity=False).hexdigest()
        return f"emb:{self.model}:{h}"

    def _get_cached(self, text: str) -> Optional[List[float]]:
        r = get_redis()
        if not r:
            return None
        try:
            val = r.get(self._cache_key(text))
            if val:
                return json.loads(val)
        except Exception:
            pass
        return None

    def _set_cache(self, text: str, emb: List[float]):
        r = get_redis()
        if not r:
            return
        try:
            r.setex(self._cache_key(text), self._cache_ttl, json.dumps(emb))
        except Exception:
            pass

    def _post(self, texts: List[str]) -> dict:
        url = f"{self.base_url}/embeddings"
        body = {"model": self.model, "input": texts}
        return self._post_embeddings(url, body)

    def _post_embeddings(self, url: str, body: dict) -> dict:
        headers = {"Content-Type": _JSON_MEDIA_TYPE}
        has_key = bool(self.api_key)
        if has_key:
            headers["Authorization"] = f"Bearer {self.api_key}"
        body_str = json.dumps(body, ensure_ascii=False)
        key_masked = ""
        if has_key:
            key = self.api_key
            key_masked = key[:8] + "..." + key[-4:] if len(key) > 12 else key
        logger.info(
            "[AI_REQ] Embedding\ncurl -X POST '%s' \\\n  -H 'Authorization: Bearer %s' \\\n  -H 'Content-Type: application/json' \\\n  -d '%s'",
            url, key_masked or "", body_str
        )
        client = _get_http_client()
        response = client.post(url, headers=headers, json=body)
        logger.info("embedding response %s %s", response.status_code, response.text[:500])
        if response.status_code != 200:
            raise RuntimeError(f"Embedding API {response.status_code}: {response.text[:300]}")
        data = response.json()
        if "error" in data:
            raise RuntimeError(f"Embedding API error: {data['error']}")
        return data

    def embed_image(self, image_b64: str, text: str = "") -> List[float]:
        """多模态图片+文本向量化（Qwen3-VL-Embedding 格式），返回图片向量。

        首选 {'image': <base64>, 'text': <text>}，若代理不识别则回退为 data URL 形式。
        """
        if not image_b64:
            raise ValueError("embed_image requires base64 image data")
        url = f"{self.base_url}/embeddings"
        attempt = [
            {"image": image_b64, "text": text or ""},
            {"image": f"data:image/jpeg;base64,{image_b64}", "text": text or ""},
        ]
        last_err = None
        for item in attempt:
            try:
                data = self._post_embeddings(url, {"model": self.model, "input": [item]})
                if data.get("data") and data["data"][0].get("embedding"):
                    return data["data"][0]["embedding"]
                last_err = f"Unexpected response: {json.dumps(data, ensure_ascii=False)[:300]}"
            except Exception as e:
                last_err = str(e)
                logger.warning("embed_image attempt failed (form=%s): %s", item.get("image")[:20], e)
        raise ValueError(f"embed_image failed: {last_err}")

    def embed(self, text: str) -> List[float]:
        cached = self._memory_cache.get(text)
        if cached is not None:
            return cached
        cached = self._get_cached(text)
        if cached is not None:
            self._memory_cache[text] = cached
            return cached
        with self._lock:
            cached = self._memory_cache.get(text)
            if cached is not None:
                return cached
            cached = self._get_cached(text)
            if cached is not None:
                self._memory_cache[text] = cached
                return cached
            data = self._post([text])
            if not data.get("data") or not data["data"][0].get("embedding"):
                raise ValueError(f"Unexpected response: {json.dumps(data, ensure_ascii=False)[:300]}")
            emb = data["data"][0]["embedding"]
            self._memory_cache[text] = emb
            self._set_cache(text, emb)
            return emb

    def embed_batch(self, texts: List[str]) -> List[List[float]]:
        if not texts:
            return []
        uncached = []
        cached_map = {}
        for t in texts:
            # check memory cache first
            c = self._memory_cache.get(t)
            if c is not None:
                cached_map[t] = c
                continue
            c = self._get_cached(t)
            if c is not None:
                self._memory_cache[t] = c
                cached_map[t] = c
            else:
                uncached.append(t)
        results = []
        if uncached:
            data = self._post(uncached)
            if not data.get("data"):
                raise ValueError(f"Empty data in response: {json.dumps(data, ensure_ascii=False)[:300]}")
            items = sorted(data["data"], key=lambda x: x.get("index", 0))
            for t, item in zip(uncached, items):
                emb = item["embedding"]
                self._memory_cache[t] = emb
                self._set_cache(t, emb)
                cached_map[t] = emb
        for t in texts:
            results.append(cached_map[t])
        return results


# ═══════════════════════════════════════════════════════════════
#  1.1 重排客户端（Reranker）
# ═══════════════════════════════════════════════════════════════
class RerankerClient:
    def __init__(self, config):
        self.config = config

    @property
    def api_key(self) -> str:
        val = self.config.reranker_api_key or self.config.llm_api_key or ""
        return val.strip()

    @property
    def base_url(self) -> str:
        raw = self.config.reranker_base_url or self.config.llm_base_url or ""
        return raw.rstrip("/")

    @property
    def model(self) -> str:
        return self.config.reranker_model or ""

    @property
    def enabled(self) -> bool:
        return self.config.reranker_enabled and bool(self.model)

    def rerank(self, query: str, documents: List[str], top_n: Optional[int] = None) -> List[Dict]:
        if not self.enabled or not documents:
            return []
        url = f"{self.base_url}/rerank"
        headers = {"Content-Type": _JSON_MEDIA_TYPE}
        has_key = bool(self.api_key)
        if has_key:
            headers["Authorization"] = f"Bearer {self.api_key}"
        body = {"model": self.model, "query": query, "documents": documents}
        if top_n is not None:
            body["top_n"] = top_n
        body_str = json.dumps(body, ensure_ascii=False)
        key_masked = ""
        if has_key:
            key = self.api_key
            key_masked = key[:8] + "..." + key[-4:] if len(key) > 12 else key
        logger.info(
            "[AI_REQ] Rerank\ncurl -X POST '%s' \\\n  -H 'Authorization: Bearer %s' \\\n  -H 'Content-Type: application/json' \\\n  -d '%s'",
            url, key_masked or "", body_str
        )
        client = _get_http_client()
        try:
            response = client.post(url, headers=headers, json=body)
            if response.status_code != 200:
                logger.warning("rerank error %s: %s", response.status_code, response.text[:300])
                return []
            data = response.json()
            results = data.get("results") or []
            return sorted(results, key=lambda x: x.get("relevance_score", 0), reverse=True)
        except Exception as e:
            logger.warning("rerank failed: %s", e)
            return []


# ═══════════════════════════════════════════════════════════════
#  1.2 VL 图片描述客户端（可选，VL 模型配置后自动启用）
# ═══════════════════════════════════════════════════════════════
class VLCaptioner:
    """用视觉对话模型为图片生成一句中文描述；未配置 VL 模型时不可用。"""

    def __init__(self, config):
        self.config = config
        self._client = None
        self._init_error = None

    @property
    def available(self) -> bool:
        return bool(self.config.vl_base_url and (self.config.vl_model or self.config.llm_model))

    def _get_client(self):
        if self._client is None and not self._init_error:
            try:
                from openai import OpenAI
                self._client = OpenAI(
                    api_key=self.config.vl_api_key or self.config.llm_api_key or "",
                    base_url=(self.config.vl_base_url or "").rstrip("/") + "/",
                )
            except Exception as e:
                self._init_error = str(e)
                logger.warning("VL client init failed: %s", e)
        return self._client

    def _log_request(self, kwargs: Dict, label: str = "VL"):
        """以 curl 命令形式记录模型请求；日志中 api-key 缩略，但真实调用使用完整 key。"""
        url = f"{(self.config.vl_base_url or self.config.llm_base_url or '').rstrip('/')}/chat/completions"
        body = json.dumps(kwargs, ensure_ascii=False, default=str)
        if len(body) > 2000:
            body = body[:2000] + "...(truncated)"
        key = self.config.vl_api_key or self.config.llm_api_key or ""
        masked = key[:8] + "..." + key[-4:] if len(key) > 12 else key
        lines = [
            f"[AI_REQ] {label}",
            f"curl -X POST '{url}' \\",
        ]
        if key:
            lines.append(f"  -H 'Authorization: Bearer {masked}' \\")
        lines.append("  -H 'Content-Type: application/json' \\")
        lines.append(f"  -d '{body}'")
        logger.info("\n".join(lines))

    def caption(self, image_b64: str, mime: str = "image/jpeg") -> str:
        if not self.available or not image_b64:
            return ""
        client = self._get_client()
        if not client:
            return ""
        model = self.config.vl_model or self.config.llm_model
        data_url = f"data:{mime};base64,{image_b64}"
        kwargs = dict(
            model=model,
            messages=[{
                "role": "user",
                "content": [
                    {"type": "text", "text": "请用一句中文描述这张图片的主要内容，30字以内，直接输出描述，不要任何前缀或解释。"},
                    {"type": "image_url", "image_url": {"url": data_url}},
                ],
            }],
            max_tokens=self.config.vl_max_tokens or 128,
            temperature=0.3,
        )
        self._log_request(kwargs, "VL caption")
        try:
            resp = client.chat.completions.create(**kwargs)
            caption = (resp.choices[0].message.content or "").strip()
            if caption:
                return caption[:200]
        except Exception as e:
            logger.warning("VL caption failed: %s", e)
        return ""


# ═══════════════════════════════════════════════════════════════
#  1.3 媒体文件解析（图片/附件文档 → 本地文件 → bytes）
# ═══════════════════════════════════════════════════════════════

def _resolve_upload_dir(upload_path: str) -> str:
    """解析上传目录：绝对路径直接用；相对路径以项目根目录为基准。"""
    p = Path(upload_path or "uploads")
    if not p.is_absolute():
        p = Path(__file__).resolve().parent.parent / p
    return str(p)


def resolve_upload_temp_path(upload_path: str) -> str:
    """AI 对话附件临时目录：解析后的 UPLOAD_PATH + "/temp"。"""
    return os.path.join(_resolve_upload_dir(upload_path or "uploads"), "temp")


def _is_remote_url(value: str) -> bool:
    """判断是否为 http/https 远端 URL（用 urlparse 避免硬编码明文协议字符串）。"""
    return urlparse(value).scheme.lower() in ("http", "https")


def _within_root(path: Path, root: Path) -> bool:
    """path（已 resolve）必须位于 root（已 resolve）之内，防止路径穿越。"""
    try:
        resolved = path.resolve()
        base = root.resolve()
    except (OSError, ValueError):
        return False
    return resolved == base or base in resolved.parents


def _local_media_path(u: str, upload_dir: str) -> Optional[str]:
    """把媒体引用解析为上传根目录内的本地路径；越界（含绝对路径/../）一律拒绝。"""
    if _is_remote_url(u):
        return None
    root = Path(upload_dir)
    if u.startswith("/files/"):
        rel = u[len("/files/"):]
        candidate = root / rel.replace("/", os.sep)
    elif u.startswith("/uploads/"):
        rel = u[len("/uploads/"):]
        candidate = root / rel.replace("/", os.sep)
    elif os.path.isabs(u):
        # 绝对路径仅在位于上传根目录内时才允许，避免读取任意文件
        candidate = Path(u)
    else:
        candidate = root / u.replace("/", os.sep)
    if not _within_root(candidate, root):
        logger.warning("blocked local media path outside upload dir: %s", u)
        return None
    return str(candidate)


def _read_local_file(path: Optional[str]) -> Optional[bytes]:
    if not (path and os.path.isfile(path)):
        return None
    try:
        with open(path, "rb") as f:
            return f.read()
    except Exception as e:
        logger.warning("read media file failed %s: %s", path, e)
        return None


def read_local_temp_media(url_or_path: str, upload_path: str = "") -> Optional[bytes]:
    """显式从 UPLOAD_PATH/temp 读取 AI 对话附件（不依赖 /uploads/ URL 前缀映射）。

    兼容 /uploads/temp/xxx、/temp/xxx、uploads/temp/xxx、纯文件名等引用形式；
    仅按文件名在上传临时目录内查找，并做越界校验，防止路径穿越。
    """
    if not url_or_path:
        return None
    u = url_or_path.strip().split("?")[0].split("#")[0].replace("\\", "/")
    if not u:
        return None
    root = Path(resolve_upload_temp_path(upload_path))
    candidate = root / u.rsplit("/", 1)[-1]
    if not _within_root(candidate, root):
        logger.warning("blocked local temp media path outside upload temp dir: %s", url_or_path)
        return None
    return _read_local_file(str(candidate))


def _allowed_media_hosts(backend_base_url: str) -> set:
    """允许抓取远端媒体的主机白名单：默认仅后端自身。"""
    hosts = set()
    if backend_base_url:
        try:
            host = (urlparse(backend_base_url).hostname or "").lower()
            if host:
                hosts.add(host)
        except ValueError:
            pass
    return hosts


def _fetch_remote_media(u: str, backend_base_url: str) -> Optional[bytes]:
    if _is_remote_url(u):
        target = u
    else:
        target = (backend_base_url or "").rstrip("/") + u
    if not _is_remote_url(target):
        return None
    host = (urlparse(target).hostname or "").lower()
    if host not in _allowed_media_hosts(backend_base_url):
        logger.warning("blocked remote media fetch to disallowed host: %s", host)
        return None
    try:
        # 禁止跟随重定向，避免白名单主机 302 到内网/云元数据造成 SSRF
        resp = _get_http_client().get(target, timeout=30, follow_redirects=False)
        if resp.status_code == 200:
            return resp.content
    except Exception as e:
        logger.warning("media file http fallback failed for %s: %s", u, e)
    return None


def resolve_media_file(url_or_path: str, upload_path: str = "", backend_base_url: str = "") -> Optional[bytes]:
    """把图片/文档的 URL（/files/...、/uploads/...、http(s)://...、本地路径）解析为文件字节。

    优先读本地 uploads 目录；本地缺失时尝试通过后端 HTTP 兜底。
    """
    if not url_or_path:
        return None
    u = url_or_path.strip().split("?")[0].split("#")[0]
    if not u:
        return None

    data = _read_local_file(_local_media_path(u, _resolve_upload_dir(upload_path)))
    if data is not None:
        return data
    return _fetch_remote_media(u, backend_base_url)


# ═══════════════════════════════════════════════════════════════
#  2. 文档解析工具
# ═══════════════════════════════════════════════════════════════
# ═══════════════════════════════════════════════════════════════
#  2.1 文本预处理规则
# ═══════════════════════════════════════════════════════════════

def preprocess_text(text: str, rules: List[str]) -> str:
    if not text or not rules:
        return text
    for rule in rules:
        if rule == "trim_whitespace":
            text = re.sub(r"[ \t\n\r]+", " ", text)
        elif rule == "remove_urls_emails":
            text = re.sub(r"https?://\S+|www\.\S+", "", text)
            text = re.sub(r"\S+@\S+\.\S+", "", text)
    return text.strip()


def parse_chunk_separator(sep: str) -> str:
    if not sep:
        return "\n\n"
    return sep.replace("\\n", "\n").replace("\\t", "\t")


# ═══════════════════════════════════════════════════════════════

def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        return _extract_pdf(file_bytes)
    elif ext == ".docx":
        return _extract_docx(file_bytes)
    elif ext in (".txt", ".md", ".csv", ".json", ".yaml", ".yml"):
        return file_bytes.decode("utf-8", errors="replace")
    else:
        return file_bytes.decode("utf-8", errors="replace")


def _extract_pdf(file_bytes: bytes) -> str:
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(io.BytesIO(file_bytes))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as e:
        logger.warning("PDF extract failed: %s", e)
        return ""


def _extract_docx(file_bytes: bytes) -> str:
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        return "\n".join(p.text for p in doc.paragraphs)
    except Exception as e:
        logger.warning("DOCX extract failed: %s", e)
        return ""


def chunk_text(text: str, max_chars: int = DEFAULT_CHUNK_SIZE,
               overlap: int = DEFAULT_CHUNK_OVERLAP,
               separators: Optional[List[str]] = None) -> List[str]:
    if not text:
        return []

    if separators is None:
        separators = ["\n\n", "\n", "。", "！", "？", ". ", "! ", "? ", "，", ", ", " "]

    chunks = []
    start = 0
    text_len = len(text)
    while start < text_len:
        end = start + max_chars
        if end >= text_len:
            remaining = text[start:].strip()
            if remaining:
                chunks.append(remaining)
            break
        boundary = _find_boundary(text, end, separators)
        chunk = text[start:boundary].strip()
        if chunk:
            chunks.append(chunk)
        next_start = boundary - overlap
        start = next_start if next_start > start else boundary
        if len(chunks) > MAX_CHUNKS_PER_DOC:
            break
    return chunks


def _find_boundary(text: str, pos: int, separators: Optional[List[str]] = None) -> int:
    if separators is None:
        separators = ["\n\n", "\n", "。", "！", "？", ". ", "! ", "? ", "，", ", ", " "]
    for sep in separators:
        idx = text.rfind(sep, pos - 100, pos)
        if idx != -1:
            return idx + len(sep)
    return pos


def estimate_tokens(text: str) -> int:
    return len(text) // 2


# ═══════════════════════════════════════════════════════════════
#  3. 知识库引擎（参考 Dify Pipeline 设计）
# ═══════════════════════════════════════════════════════════════
class KnowledgeEngine:
    def __init__(self, config):
        self.config = config
        self.embedder = EmbeddingClient(config)
        self.reranker = RerankerClient(config)
        self.vl_captioner = VLCaptioner(config)
        self._ensure_vector_column()
        self._ensure_kb_columns()
        self._ensure_chunk_doc_columns()

    # ── 数据库初始化 ─────────────────────────────────────────
    def _ensure_vector_column(self):
        """确保 knowledge_chunks 表存在 embedding vector 列（维度与当前模型匹配）"""
        dim = self.embedder.dimension
        try:
            engine = get_db_engine()
            with engine.connect() as conn:
                conn.execute(sa_text("CREATE EXTENSION IF NOT EXISTS vector"))
                conn.execute(sa_text(
                    f"ALTER TABLE knowledge_chunks ADD COLUMN IF NOT EXISTS embedding vector({dim})"
                ))
                # 若列已存在但维度不符（如 Java 初始化了 vector(2048)），强制修改
                conn.execute(sa_text(
                    f"ALTER TABLE knowledge_chunks ALTER COLUMN embedding TYPE vector({dim})"
                ))
                conn.commit()
        except Exception as e:
            logger.warning("vector column init: %s", e)

    def _ensure_kb_columns(self):
        """向 knowledge_bases 表添加新配置列（如不存在）"""
        new_cols = [
            ("chunk_size", "INTEGER DEFAULT 1024"),
            ("chunk_overlap", "INTEGER DEFAULT 50"),
            ("chunk_separator", "VARCHAR(50) DEFAULT '\n\n'"),
            ("text_preprocessing_rules", "JSON DEFAULT '[]'::json"),
            ("embedding_model", "VARCHAR(255) DEFAULT NULL"),
            ("embedding_dimension", "INTEGER DEFAULT NULL"),
            ("top_k", "INTEGER DEFAULT 6"),
            ("threshold_min", "FLOAT DEFAULT 0.5"),
            ("threshold_max", "FLOAT DEFAULT 0.7"),
            ("reranker_model", "VARCHAR(255) DEFAULT NULL"),
            ("reranker_enabled", "BOOLEAN DEFAULT FALSE"),
        ]
        engine = get_db_engine()
        with engine.connect() as conn:
            for col_name, col_type in new_cols:
                try:
                    conn.execute(sa_text(
                        f"ALTER TABLE knowledge_bases ADD COLUMN IF NOT EXISTS {col_name} {col_type}"
                    ))
                except Exception as e:
                    logger.warning("add column %s: %s", col_name, e)
            conn.commit()

    def _ensure_chunk_doc_columns(self):
        """确保 knowledge_chunks.media_type 与 knowledge_documents.images 列存在"""
        engine = get_db_engine()
        statements = [
            "ALTER TABLE knowledge_chunks ADD COLUMN IF NOT EXISTS media_type VARCHAR(20) DEFAULT 'text'",
            "ALTER TABLE knowledge_documents ADD COLUMN IF NOT EXISTS images JSON",
        ]
        with engine.connect() as conn:
            for stmt in statements:
                try:
                    conn.execute(sa_text(stmt))
                except Exception as e:
                    logger.warning("ensure chunk/doc column: %s", e)
            conn.commit()

    def ensure_tables(self):
        Base.metadata.create_all(get_db_engine())

    # ── 知识库 CRUD ──────────────────────────────────────────
    def create_kb(self, name: str, description: str = "",
                  chunk_size: int = DEFAULT_CHUNK_SIZE,
                  chunk_overlap: int = DEFAULT_CHUNK_OVERLAP,
                  chunk_separator: str = DEFAULT_CHUNK_SEPARATOR,
                  text_preprocessing_rules: Optional[List[str]] = None,
                  embedding_model: Optional[str] = None,
                  embedding_dimension: Optional[int] = None,
                  top_k: int = 6,
                  threshold_min: float = 0.5,
                  threshold_max: float = 0.7,
                  reranker_model: Optional[str] = None,
                  reranker_enabled: bool = False) -> KnowledgeBase:
        with get_db() as db:
            kb = KnowledgeBase(
                name=name, description=description,
                chunk_size=chunk_size, chunk_overlap=chunk_overlap,
                chunk_separator=chunk_separator,
                text_preprocessing_rules=text_preprocessing_rules or [],
                embedding_model=embedding_model, embedding_dimension=embedding_dimension,
                top_k=top_k, threshold_min=threshold_min, threshold_max=threshold_max,
                reranker_model=reranker_model, reranker_enabled=reranker_enabled,
            )
            db.add(kb)
            db.flush()
            db.refresh(kb)
            return kb

    def update_kb_config(self, kb_id: int, **kwargs) -> Optional[KnowledgeBase]:
        allowed_keys = {
            "name", "description", "chunk_size", "chunk_overlap", "chunk_separator",
            "text_preprocessing_rules", "embedding_model", "embedding_dimension",
            "top_k", "threshold_min", "threshold_max", "reranker_model", "reranker_enabled",
        }
        with get_db() as db:
            kb = db.execute(
                select(KnowledgeBase).where(KnowledgeBase.id == kb_id)
            ).scalar_one_or_none()
            if not kb:
                return None
            for key, val in kwargs.items():
                if key in allowed_keys and val is not None:
                    setattr(kb, key, val)
            db.flush()
            db.refresh(kb)
            return kb

    def list_kbs(self) -> List[Dict]:
        with get_db() as db:
            rows = db.execute(
                select(KnowledgeBase).order_by(KnowledgeBase.updated_at.desc())
            ).scalars().all()
            return [self._kb_to_dict(r) for r in rows]

    def get_kb(self, kb_id: int) -> Optional[KnowledgeBase]:
        with get_db() as db:
            return db.execute(
                select(KnowledgeBase).where(KnowledgeBase.id == kb_id)
            ).scalar_one_or_none()

    def delete_kb(self, kb_id: int):
        with get_db() as db:
            db.execute(sa_delete(KnowledgeBase).where(KnowledgeBase.id == kb_id))

    def _kb_to_dict(self, kb: KnowledgeBase) -> Dict:
        rules = kb.text_preprocessing_rules
        if isinstance(rules, str):
            try:
                rules = json.loads(rules)
            except (json.JSONDecodeError, TypeError):
                rules = []
        return {
            "id": kb.id,
            "name": kb.name,
            "description": _or_default(kb.description, ""),
            "status": kb.status,
            "chunk_size": _or_default(kb.chunk_size, DEFAULT_CHUNK_SIZE),
            "chunk_overlap": _or_default(kb.chunk_overlap, DEFAULT_CHUNK_OVERLAP),
            "chunk_separator": _or_default(kb.chunk_separator, DEFAULT_CHUNK_SEPARATOR),
            "text_preprocessing_rules": rules if isinstance(rules, list) else [],
            "embedding_model": _or_default(kb.embedding_model, ""),
            "embedding_dimension": _or_default(kb.embedding_dimension, 0),
            "top_k": _or_default(kb.top_k, 6),
            "threshold_min": _or_default(kb.threshold_min, 0.5),
            "threshold_max": _or_default(kb.threshold_max, 0.7),
            "reranker_model": _or_default(kb.reranker_model, ""),
            "reranker_enabled": _or_default(kb.reranker_enabled, False),
            "created_at": kb.created_at.isoformat() if kb.created_at else "",
            "updated_at": kb.updated_at.isoformat() if kb.updated_at else "",
        }

    # ── 文档管理 ─────────────────────────────────────────────
    def add_document(self, kb_id: int, title: str, content: str,
                     file_name: str = "", file_size: int = 0) -> KnowledgeDocument:
        with get_db() as db:
            doc = KnowledgeDocument(
                knowledge_base_id=kb_id,
                title=title,
                content=content,
                file_name=file_name,
                file_size=file_size,
                status=0,
            )
            db.add(doc)
            db.flush()
            db.refresh(doc)
            return doc

    def list_documents(self, kb_id: int) -> List[Dict]:
        with get_db() as db:
            rows = db.execute(
                select(KnowledgeDocument)
                .where(KnowledgeDocument.knowledge_base_id == kb_id)
                .order_by(KnowledgeDocument.updated_at.desc())
            ).scalars().all()
            return [self._doc_to_dict(r) for r in rows]

    def get_document(self, doc_id: int) -> Optional[KnowledgeDocument]:
        with get_db() as db:
            return db.execute(
                select(KnowledgeDocument).where(KnowledgeDocument.id == doc_id)
            ).scalar_one_or_none()

    def delete_document(self, doc_id: int):
        with get_db() as db:
            db.execute(sa_delete(KnowledgeDocument).where(KnowledgeDocument.id == doc_id))

    def _doc_to_dict(self, doc: KnowledgeDocument) -> Dict:
        return {
            "id": doc.id,
            "knowledge_base_id": doc.knowledge_base_id,
            "title": doc.title,
            "file_name": doc.file_name or "",
            "file_size": doc.file_size or 0,
            "status": doc.status,
            "chunk_count": doc.chunk_count or 0,
            "created_at": doc.created_at.isoformat() if doc.created_at else "",
            "updated_at": doc.updated_at.isoformat() if doc.updated_at else "",
        }

    def _get_kb_config(self, kb_id: int) -> Dict:
        kb = self.get_kb(kb_id)
        if not kb:
            return {}
        rules_raw = kb.text_preprocessing_rules
        if isinstance(rules_raw, str):
            try:
                rules_raw = json.loads(rules_raw)
            except (json.JSONDecodeError, TypeError):
                rules_raw = []
        return {
            "chunk_size": kb.chunk_size or DEFAULT_CHUNK_SIZE,
            "chunk_overlap": kb.chunk_overlap or DEFAULT_CHUNK_OVERLAP,
            "chunk_separator": kb.chunk_separator or DEFAULT_CHUNK_SEPARATOR,
            "text_preprocessing_rules": rules_raw if isinstance(rules_raw, list) else [],
            "embedding_model": kb.embedding_model or "",
            "embedding_dimension": kb.embedding_dimension or 0,
            "top_k": kb.top_k or 6,
            "threshold_min": kb.threshold_min or 0.5,
            "threshold_max": kb.threshold_max or 0.7,
            "reranker_model": kb.reranker_model or "",
            "reranker_enabled": kb.reranker_enabled or False,
        }

    # ── 文档处理管线（分块 → 嵌入 → 存储）────────────────
    def _doc_text_chunks(self, doc: KnowledgeDocument, chunk_size: int, chunk_overlap: int,
                         separators: List[str]) -> List[str]:
        if not doc.file_path:
            return []
        upload_path = getattr(self.config, "upload_path", "")
        backend_base_url = getattr(self.config, "backend_base_url", "")
        file_bytes = resolve_media_file(doc.file_path, upload_path, backend_base_url)
        if not file_bytes:
            return []
        fname = doc.file_name or os.path.basename(doc.file_path or "")
        file_text = extract_text_from_file(file_bytes, fname)
        if not file_text.strip():
            return []
        return chunk_text(file_text, max_chars=chunk_size, overlap=chunk_overlap, separators=separators)

    def _image_items(self, doc: KnowledgeDocument) -> List:
        if not doc.images:
            return []
        items = []
        for img in doc.images if isinstance(doc.images, list) else []:
            if not isinstance(img, dict):
                continue
            url = str(img.get("url") or "")
            caption = str(img.get("caption") or "")
            if not url:
                continue
            items.append((url, caption))
        return items

    def _embed_image_items(self, image_items: List, upload_path: str, backend_base_url: str) -> List:
        records = []
        for url, caption in image_items:
            try:
                file_bytes = resolve_media_file(url, upload_path, backend_base_url)
                if not file_bytes:
                    logger.warning("skip image, file not found: %s", url)
                    continue
                ext = os.path.splitext(url.split("/")[-1])[1].lower()
                mime = mimetypes.guess_type("x" + ext)[0] or "image/jpeg"
                image_b64 = base64.b64encode(file_bytes).decode("ascii")
                description = self.vl_captioner.caption(image_b64, mime) if self.vl_captioner.available else ""
                if not description:
                    description = caption or f"图片路径：{url}"
                emb = self.embedder.embed_image(image_b64, description)
                chunk_content = f"图片描述：{description}\n图片路径：{url}"
                records.append((chunk_content, emb, "image"))
            except Exception as e:
                logger.warning("image vectorize failed for %s: %s", url, e)
        return records

    def _collect_records(self, doc: KnowledgeDocument) -> List:
        kb_config = self._get_kb_config(doc.knowledge_base_id)
        upload_path = getattr(self.config, "upload_path", "")
        backend_base_url = getattr(self.config, "backend_base_url", "")
        content = doc.content
        rules = kb_config.get("text_preprocessing_rules", [])
        if rules:
            content = preprocess_text(content, rules)
        sep = parse_chunk_separator(kb_config.get("chunk_separator", DEFAULT_CHUNK_SEPARATOR))
        chunk_size = kb_config.get("chunk_size", DEFAULT_CHUNK_SIZE)
        chunk_overlap = kb_config.get("chunk_overlap", DEFAULT_CHUNK_OVERLAP)
        separators = [sep, "\n", "。", "！", "？", ". ", "! ", "? ", "，", ", ", " "]
        text_chunks = chunk_text(content, max_chars=chunk_size, overlap=chunk_overlap, separators=separators)
        doc_chunks = self._doc_text_chunks(doc, chunk_size, chunk_overlap, separators)
        image_items = self._image_items(doc)

        records = []
        if text_chunks:
            text_embeddings = self.embedder.embed_batch(text_chunks)
            for chunk, emb in zip(text_chunks, text_embeddings):
                records.append((chunk, emb, "text"))
        if doc_chunks:
            doc_embeddings = self.embedder.embed_batch(doc_chunks)
            for chunk, emb in zip(doc_chunks, doc_embeddings):
                records.append((chunk, emb, "doc"))
        records.extend(self._embed_image_items(image_items, upload_path, backend_base_url))
        return records

    def _store_chunks(self, doc: KnowledgeDocument, records: List):
        dim = self.embedder.dimension
        with get_db() as db:
            db.execute(
                sa_delete(KnowledgeChunk).where(KnowledgeChunk.document_id == doc.id)
            )
            for i, (chunk, emb, media_type) in enumerate(records):
                tokens = estimate_tokens(chunk)
                emb_str = "[" + ",".join(str(v) for v in emb[:dim]) + "]"
                db.execute(sa_text("""
                    INSERT INTO knowledge_chunks
                        (document_id, knowledge_base_id, content, media_type, chunk_index, tokens, embedding, created_at)
                    VALUES (:doc_id, :kb_id, :content, :media_type, :idx, :tokens, CAST(:embedding AS vector), :now)
                """), {
                    "doc_id": doc.id,
                    "kb_id": doc.knowledge_base_id,
                    "content": chunk,
                    "media_type": media_type,
                    "idx": i,
                    "tokens": tokens,
                    "embedding": emb_str,
                    "now": _utcnow(),
                })

            db.execute(
                sa_text("UPDATE knowledge_documents SET status = 1, chunk_count = :cnt, updated_at = :now WHERE id = :id"),
                {"cnt": len(records), "id": doc.id, "now": _utcnow()},
            )

    def process_document(self, doc_id: int) -> Dict:
        doc = self.get_document(doc_id)
        if not doc:
            return {"status": "error", "message": "Document not found"}

        try:
            records = self._collect_records(doc)

            if not records:
                doc.status = 3  # error
                return {"status": "error", "message": "Empty content after chunking"}

            self._store_chunks(doc, records)
            return {"status": "ok", "chunks": len(records), "document_id": doc.id}

        except Exception as e:
            logger.exception("process_document error")
            with get_db() as db:
                db.execute(
                    sa_text("UPDATE knowledge_documents SET status = 3 WHERE id = :id"),
                    {"id": doc.id},
                )
            return {"status": "error", "message": str(e)}

    # ── 向量搜索 ─────────────────────────────────────────────
    def search(self, kb_id: int, query: str, top_k: Optional[int] = None,
               threshold: Optional[float] = None, max_threshold: Optional[float] = None) -> List[Dict]:
        kb_config = self._get_kb_config(kb_id)
        top_k = top_k if top_k is not None else kb_config.get("top_k", 6)
        threshold = threshold if threshold is not None else kb_config.get("threshold_min", 0.5)
        max_threshold = max_threshold if max_threshold is not None else 1.0
        query_emb = self.embedder.embed(query)
        dim = self.embedder.dimension
        emb_str = "[" + ",".join(str(v) for v in query_emb[:dim]) + "]"

        # max_threshold 为 None 或 >=1.0 时去掉上限（避免带通滤掉最高相似度结果）
        similarity_expr = "1 - (kc.embedding <=> CAST(:query_emb AS vector))"
        if max_threshold is None or max_threshold >= 1.0:
            range_cond = f"{similarity_expr} >= :threshold"
            params = {"query_emb": emb_str, "kb_id": kb_id, "top_k": top_k, "threshold": threshold}
        else:
            range_cond = f"{similarity_expr} BETWEEN :threshold AND :max_threshold"
            params = {"query_emb": emb_str, "kb_id": kb_id, "top_k": top_k,
                      "threshold": threshold, "max_threshold": max_threshold}

        with get_db() as db:
            rows = db.execute(sa_text(f"""
                SELECT
                    kc.id,
                    kc.content,
                    kc.media_type,
                    kc.chunk_index,
                    kc.tokens,
                    kc.document_id,
                    kd.title as doc_title,
                    {similarity_expr} AS similarity
                FROM knowledge_chunks kc
                JOIN knowledge_documents kd ON kd.id = kc.document_id
                WHERE kc.knowledge_base_id = :kb_id
                  AND kc.embedding IS NOT NULL
                  AND {range_cond}
                ORDER BY {similarity_expr}
                LIMIT :top_k
            """), params).fetchall()

            results = []
            for r in rows:
                results.append({
                    "id": r[0],
                    "content": r[1],
                    "media_type": r[2] or "text",
                    "chunk_index": r[3],
                    "tokens": r[4],
                    "document_id": r[5],
                    "doc_title": r[6],
                    "similarity": round(float(r[7]), 4),
                })
            return results

    def search_all_kbs(self, query: str, top_k: int = 5,
                       threshold: float = 0.0, max_threshold: float = 1.0) -> List[Dict]:
        query_emb = self.embedder.embed(query)
        dim = self.embedder.dimension
        emb_str = "[" + ",".join(str(v) for v in query_emb[:dim]) + "]"

        similarity_expr = "1 - (kc.embedding <=> CAST(:query_emb AS vector))"
        if max_threshold is None or max_threshold >= 1.0:
            range_cond = f"{similarity_expr} >= :threshold"
            params = {"query_emb": emb_str, "top_k": top_k, "threshold": threshold}
        else:
            range_cond = f"{similarity_expr} BETWEEN :threshold AND :max_threshold"
            params = {"query_emb": emb_str, "top_k": top_k,
                      "threshold": threshold, "max_threshold": max_threshold}

        with get_db() as db:
            rows = db.execute(sa_text(f"""
                SELECT
                    kc.id,
                    kc.content,
                    kc.media_type,
                    kc.chunk_index,
                    kc.tokens,
                    kc.document_id,
                    kd.title as doc_title,
                    kb.name as kb_name,
                    kc.knowledge_base_id,
                    {similarity_expr} AS similarity
                FROM knowledge_chunks kc
                JOIN knowledge_documents kd ON kd.id = kc.document_id
                JOIN knowledge_bases kb ON kb.id = kc.knowledge_base_id
                WHERE kc.embedding IS NOT NULL
                  AND {range_cond}
                ORDER BY {similarity_expr}
                LIMIT :top_k
            """), params).fetchall()

            results = []
            for r in rows:
                results.append({
                    "id": r[0],
                    "content": r[1],
                    "media_type": r[2] or "text",
                    "chunk_index": r[3],
                    "tokens": r[4],
                    "document_id": r[5],
                    "doc_title": r[6],
                    "kb_name": r[7],
                    "knowledge_base_id": r[8],
                    "similarity": round(float(r[9]), 4),
                })
            return results

    # ── RAG 上下文构建 ────────────────────────────────────────
    def _gather_rag_results(self, query: str, kb_ids: Optional[List[int]],
                            top_k: Optional[int], threshold: Optional[float],
                            max_threshold: Optional[float]) -> List[Dict]:
        if kb_ids:
            results = []
            for kb_id in kb_ids:
                results.extend(self.search(kb_id, query, top_k, threshold, max_threshold))
            results.sort(key=lambda x: x["similarity"], reverse=True)
            return results[:top_k or 6]
        return self.search_all_kbs(query, top_k or 6, threshold or 0.0, max_threshold or 1.0)

    def _maybe_rerank_results(self, query: str, results: List[Dict], top_k: Optional[int],
                              use_reranker: Optional[bool]) -> List[Dict]:
        do_rerank = use_reranker if use_reranker is not None else self.reranker.enabled
        if not (do_rerank and len(results) > 1):
            return results
        docs = [r["content"] for r in results]
        reranked = self.reranker.rerank(query, docs, top_n=top_k or 6)
        if not reranked:
            return results
        reranked_results = []
        seen = set()
        for rr in reranked:
            idx = rr.get("index")
            if idx is not None and idx < len(results) and idx not in seen:
                r = dict(results[idx])
                r["similarity"] = rr.get("relevance_score", r["similarity"])
                reranked_results.append(r)
                seen.add(idx)
        return reranked_results or results

    def _format_rag_context(self, results: List[Dict]) -> str:
        parts = []
        for r in results:
            source = r.get("doc_title", "") or f"Doc#{r['document_id']}"
            parts.append(f"[来源: {source} (相似度: {r['similarity']:.2f})]\n{r['content']}")
        return "\n\n---\n\n".join(parts)

    def build_rag_context(self, query: str, kb_ids: Optional[List[int]] = None,
                          top_k: Optional[int] = None, threshold: Optional[float] = None,
                          max_threshold: Optional[float] = None,
                          use_reranker: Optional[bool] = None) -> str:
        results = self._gather_rag_results(query, kb_ids, top_k, threshold, max_threshold)
        if not results:
            return ""
        results = self._maybe_rerank_results(query, results, top_k, use_reranker)
        return self._format_rag_context(results)

    # ── 通过 source_type/source_id 来同步文档（upsert）───
    def sync_document(self, kb_id: int, source_type: str, source_id: int,
                      title: str, content: str,
                      file_path: str = "", images: Optional[List[Dict]] = None) -> Dict:
        with get_db() as db:
            existing = db.execute(
                select(KnowledgeDocument)
                .where(KnowledgeDocument.source_type == source_type)
                .where(KnowledgeDocument.source_id == source_id)
            ).scalar_one_or_none()

            if existing:
                existing.title = title
                existing.content = content
                if file_path:
                    existing.file_path = file_path
                if images is not None:
                    existing.images = images or None
                existing.status = 0
                db.flush()
                doc_id = existing.id
            else:
                doc = KnowledgeDocument(
                    knowledge_base_id=kb_id,
                    title=title,
                    content=content,
                    file_path=file_path or "",
                    images=images or None,
                    source_type=source_type,
                    source_id=source_id,
                    status=0,
                )
                db.add(doc)
                db.flush()
                db.refresh(doc)
                doc_id = doc.id

        result = self.process_document(doc_id)
        return {"document_id": doc_id, "status": result.get("status"), "chunks": result.get("chunks", 0)}

    # ── 通过 source_type/source_id 删除文档 ─────────────────
    def delete_document_by_source(self, source_type: str, source_id: int) -> bool:
        with get_db() as db:
            doc = db.execute(
                select(KnowledgeDocument)
                .where(KnowledgeDocument.source_type == source_type)
                .where(KnowledgeDocument.source_id == source_id)
            ).scalar_one_or_none()
            if not doc:
                return False
            db.delete(doc)
            return True

    # ── 重嵌入所有块 ─────────────────────────────────────────
    def _load_chunks_for_reembed(self, kb_id: Optional[int]):
        with get_db() as db:
            query = select(KnowledgeChunk)
            if kb_id:
                query = query.where(KnowledgeChunk.knowledge_base_id == kb_id)
            return db.execute(query).scalars().all()

    def _reembed_text_chunks(self, chunks) -> Dict:
        text_chunks = [c for c in chunks if (c.media_type or "text") in ("text", "doc")]
        text_embeddings = self.embedder.embed_batch([c.content for c in text_chunks]) if text_chunks else []
        return {c.id: emb for c, emb in zip(text_chunks, text_embeddings)}

    def _reembed_image_chunks(self, chunks, upload_path: str, backend_base_url: str) -> Dict:
        img_map = {}
        for c in chunks:
            if (c.media_type or "text") != "image":
                continue
            m = re.search(r"图片路径：(\S+)", c.content or "")
            url = m.group(1) if m else ""
            try:
                file_bytes = resolve_media_file(url, upload_path, backend_base_url) if url else None
                if not file_bytes:
                    continue
                description = (c.content or "").replace(f"图片路径：{url}", "").replace("图片描述：", "").strip()
                image_b64 = base64.b64encode(file_bytes).decode("ascii")
                img_map[c.id] = self.embedder.embed_image(image_b64, description)
            except Exception as e:
                logger.warning("reembed image chunk %s failed: %s", c.id, e)
        return img_map

    def _update_chunk_embeddings(self, chunks, text_map: Dict, img_map: Dict, dim: int) -> int:
        count = 0
        with get_db() as db:
            for c in chunks:
                emb = text_map.get(c.id) or img_map.get(c.id)
                if emb is None:
                    continue
                emb_str = "[" + ",".join(str(v) for v in emb[:dim]) + "]"
                db.execute(sa_text("""
                    UPDATE knowledge_chunks SET embedding = CAST(:emb AS vector) WHERE id = :id
                """), {"emb": emb_str, "id": c.id})
                count += 1
        return count

    def reembed_all(self, kb_id: Optional[int] = None) -> Dict:
        chunks = self._load_chunks_for_reembed(kb_id)
        if not chunks:
            return {"status": "ok", "reembedded": 0}

        upload_path = getattr(self.config, "upload_path", "")
        backend_base_url = getattr(self.config, "backend_base_url", "")
        dim = self.embedder.dimension
        text_map = self._reembed_text_chunks(chunks)
        img_map = self._reembed_image_chunks(chunks, upload_path, backend_base_url)
        count = self._update_chunk_embeddings(chunks, text_map, img_map, dim)
        return {"status": "ok", "reembedded": count}
